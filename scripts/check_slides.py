from pathlib import Path
import json
import fitz
from PIL import Image,ImageDraw
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
pdf=fitz.open(ROOT/'slides/论文复现汇报.pdf')
prs=Presentation(ROOT/'slides/论文复现汇报.pptx')
assert len(pdf)==len(prs.slides)==24
notes=json.loads((ROOT/'slides/slide_plan.json').read_text())
assert sum(n['seconds'] for n in notes if not n['backup'])==1500
assert sum(not n['backup'] for n in notes)==20
assert [n['slide'] for n in notes if n['backup']]==[21,22,23,24]
assert sum(n['seconds'] for n in notes if n['backup'])==300
for index,label in [(20,'问题一'),(21,'答案一'),(22,'问题二'),(23,'答案二')]:
    content=' '.join(sh.text for sh in prs.slides[index].shapes if sh.has_text_frame)
    assert label in content,(index,content)
    if index in [20,22]: assert '答案一：' not in content and '答案二：' not in content
for n in notes:
    assert n['script'] in prs.slides[n['slide']-1].notes_slide.notes_text_frame.text
for i,sl in enumerate(prs.slides):
    assert sl.has_notes_slide
    for sh in sl.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=prs.slide_width+1000,(i,sh.name)
        assert sh.top+sh.height<=prs.slide_height+1000,(i,sh.name)
    assert len(pdf[i].get_text())>70,(i,'Missing slide text')
for group in range(3):
    pages=list(range(group*10,min(group*10+10,len(pdf))))
    sheet=Image.new('RGB',(1600,5*480),'#dfe6eb');draw=ImageDraw.Draw(sheet)
    for j,i in enumerate(pages):
        pix=pdf[i].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False)
        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        im.thumbnail((780,439));x=(j%2)*800+10;y=(j//2)*480+25
        sheet.paste(im,(x,y));draw.text((x,y-20),f'Slide {i+1}',fill='black')
    sheet.save(ROOT/f'slides/contact_sheet_{group+1}.png')
# 重点页面保留大图，用于字体/图表可读性检查。
for i in [0,2,4,7,9,14,15,17]:
    pix=pdf[i].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
    pix.save(ROOT/f'slides/preview_{i+1:02d}.png')
print('PASS: 24 slides; 20 main + 2 question/answer pairs; notes present; main 25:00 + Q&A 5:00')
