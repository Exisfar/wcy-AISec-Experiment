from pathlib import Path
import json
import fitz
from PIL import Image,ImageDraw
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
pdf=fitz.open(ROOT/'presentation/论文复现汇报.pdf')
prs=Presentation(ROOT/'presentation/论文复现汇报.pptx')
assert len(pdf)==len(prs.slides)==25
notes=json.loads((ROOT/'presentation/slide_plan.json').read_text())
assert sum(n['seconds'] for n in notes if not n['backup'])==1500
assert sum(not n['backup'] for n in notes)==21
assert [n['slide'] for n in notes if n['backup']]==[22,23,24,25]
assert sum(n['seconds'] for n in notes if n['backup'])==300
audience_text='\n'.join(sh.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame)
audience_text+='\n'+'\n'.join(n['script'] for n in notes)
for phrase in ['文件名','连字符仓库','串题','Kaggle版','Kaggle 版','从Kaggle下载的压缩数据','从Kaggle下载的压缩数据','Kaggle改动版','Kaggle 改动版','对照评分要求']:
    assert phrase not in audience_text,('课堂材料出现不合适的过程叙述',phrase)
assert notes[1]['title']=='汇报目录'
agenda_text=' '.join(sh.text for sh in prs.slides[1].shapes if sh.has_text_frame)
assert all(label in agenda_text for label in ['论文内容','复现工作','不足与改进','代码演示','讨论'])
assert '实验结果' not in agenda_text
stat_text=' '.join(sh.text for sh in prs.slides[11].shapes if sh.has_text_frame)
assert '独立实现指标公式' not in stat_text
assert '作者开源实现' in stat_text
assert '05' in agenda_text
assert any(sh.has_text_frame and sh.text=='代码演示' for sh in prs.slides[19].shapes)
assert '总体均衡' in notes[2]['title']
assert '论文' in notes[3]['title']
demo_text='\n'.join(sh.text for sh in prs.slides[19].shapes if sh.has_text_frame)
assert 'bias-repro-gpu' in demo_text and '/workspace/notebooks' in demo_text
assert 'docker run' not in audience_text and 'reproduce 命令' not in audience_text
for number in ['01','02','03']:assert number in demo_text
assert '00' not in demo_text
assert '四份' not in demo_text and '三份' in demo_text
for index,label in [(21,'问题一'),(22,'答案一'),(23,'问题二'),(24,'答案二')]:
    content=' '.join(sh.text for sh in prs.slides[index].shapes if sh.has_text_frame)
    assert label in content,(index,content)
    if index in [21,23]: assert '答案一：' not in content and '答案二：' not in content
for n in notes:
    assert n['script'] == prs.slides[n['slide']-1].notes_slide.notes_text_frame.text
for i,sl in enumerate(prs.slides):
    assert sl.has_notes_slide
    for shape in sl.shapes:
        if shape.has_text_frame and shape.text.strip():
            for paragraph in shape.text_frame.paragraphs:
                assert paragraph.font.name == 'Microsoft YaHei'
                for run in paragraph.runs:
                    props=run._r.get_or_add_rPr()
                    assert props.get('lang') == 'zh-CN'
                    assert props.find('a:ea',props.nsmap).get('typeface') == 'Microsoft YaHei'

    for sh in sl.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=prs.slide_width+1000,(i,sh.name)
        assert sh.top+sh.height<=prs.slide_height+1000,(i,sh.name)
    pdf_text=''.join(pdf[i].get_text().split())
    assert len(pdf_text)>30,(i,'Missing slide text')
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text.strip():
            expected=''.join(sh.text.split())
            assert expected in pdf_text,(i,'PDF text differs from PPT',sh.text[:60])
# 检查不同文本框的交叠，避免新增脚注或说明压入表格。
for page_index,slide in enumerate(prs.slides,1):
    boxes=[shape for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
    for position,a in enumerate(boxes):
        for b in boxes[position+1:]:
            width=min(a.left+a.width,b.left+b.width)-max(a.left,b.left)
            height=min(a.top+a.height,b.top+b.height)-max(a.top,b.top)
            assert not (width>20000 and height>20000),(page_index,'Text overlap',a.text[:35],b.text[:35])
# 公式和曲线以图片插入；同时检查图片与外部文字是否相交。
for page_index,slide in enumerate(prs.slides,1):
    if page_index==1:continue  # 封面参考背景上叠加标题为预期布局。
    objects=[shape for shape in slide.shapes if shape.shape_type==13 or (shape.has_text_frame and shape.text.strip())]
    for position,a in enumerate(objects):
        for b in objects[position+1:]:
            width=min(a.left+a.width,b.left+b.width)-max(a.left,b.left)
            height=min(a.top+a.height,b.top+b.height)-max(a.top,b.top)
            assert not (width>20000 and height>20000),(page_index,'Content overlap',a.name,b.name)
for page_index,minimum in [(4,6),(10,2),(12,2),(17,3)]:
    assert sum(shape.shape_type==13 for shape in prs.slides[page_index].shapes)>=minimum,('Missing equations',page_index+1)
for group in range(3):
    pages=list(range(group*10,min(group*10+10,len(pdf))))
    sheet=Image.new('RGB',(1600,5*480),'#dfe6eb');draw=ImageDraw.Draw(sheet)
    for j,i in enumerate(pages):
        pix=pdf[i].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False)
        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        im.thumbnail((780,439));x=(j%2)*800+10;y=(j//2)*480+25
        sheet.paste(im,(x,y));draw.text((x,y-20),f'Slide {i+1}',fill='black')
    sheet.save(ROOT/f'presentation/contact_sheet_{group+1}.png')
# 重点页面保留大图，用于字体/图表可读性检查。
for i in [0,1,2,3,4,5,8,10,12,13,14,15,16,17,18,19]:
    pix=pdf[i].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
    pix.save(ROOT/f'presentation/preview_{i+1:02d}.png')
print('PASS: 25 slides; 21 main + 2 question/answer pairs; notes present; main 25:00 + Q&A 5:00')
